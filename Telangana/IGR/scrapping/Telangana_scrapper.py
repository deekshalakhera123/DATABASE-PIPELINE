import os
import time
import threading
import pandas as pd
from openpyxl import load_workbook
from selenium import webdriver
from selenium.common.exceptions import TimeoutException, NoSuchWindowException
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from PIL import Image, ImageOps, ImageFilter
import pytesseract
import requests
from queue import Queue

pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

excel_file = r"D:\TELANGANA RERA\claudeCode\registration_details_2025.xlsx"
progress_file = "last_doc.txt"

# ─── Thread-safe locks ────────────────────────────────────────────────────────
excel_lock = threading.Lock()
progress_lock = threading.Lock()

NUM_THREADS = 2

# ─── Progress / doc-list reader ───────────────────────────────────────────────
def read_doc_list():
    """
    Reads last_doc.txt and returns (mode, data):

    COMMA mode  → file contains:  4911, 4912, 4913, 4914
                  returns ('list', [4911, 4912, 4913, 4914])

    RANGE mode  → file contains a single integer (last completed doc)
                  returns ('range', 4911)   ← resume from that number

    MISSING     → file absent, start fresh from 1
                  returns ('range', 1)
    """
    if not os.path.exists(progress_file):
        return ('range', 1)

    with open(progress_file, 'r') as f:
        content = f.read().strip()

    if ',' in content:
        # Comma-separated list of specific doc numbers
        doc_numbers = [int(x.strip()) for x in content.split(',') if x.strip().isdigit()]
        return ('list', doc_numbers)
    else:
        # Single number — sequential range resume point
        return ('range', int(content))

def save_progress(doc_no):
    """Only used in range mode to track last completed doc."""
    with progress_lock:
        with open(progress_file, 'w') as f:
            f.write(str(doc_no))

# ─── Driver setup ─────────────────────────────────────────────────────────────
def setup_driver():
    options = webdriver.ChromeOptions()
    #options.add_argument("--headless=new")
    options.add_argument("--disable-gpu")
    options.add_argument("--window-size=1920,1080")
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    driver = webdriver.Chrome(options=options)
    return driver

# ─── CAPTCHA solver ───────────────────────────────────────────────────────────
def download_and_solve_captcha(driver):
    captcha_element = WebDriverWait(driver, 10).until(
        EC.presence_of_element_located((By.XPATH, '//*[@id="myForm"]/div[4]/img'))
    )
    captcha_url = captcha_element.get_attribute("src")
    cookies = driver.get_cookies()
    session = requests.Session()
    for cookie in cookies:
        session.cookies.set(cookie['name'], cookie['value'])

    response = session.get(captcha_url, stream=True)
    captcha_path = f"captcha_{threading.current_thread().name}.png"
    with open(captcha_path, 'wb') as f:
        f.write(response.content)

    img = Image.open(captcha_path).convert("L")
    img = ImageOps.autocontrast(img)
    img = img.point(lambda x: 0 if x < 150 else 255)
    img = img.filter(ImageFilter.MedianFilter(size=3))
    img = img.resize((img.width * 2, img.height * 2), Image.Resampling.LANCZOS)

    text = pytesseract.image_to_string(
        img, config='--psm 7 -c tessedit_char_whitelist=ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789'
    ).strip()
    return text

# ─── Login ────────────────────────────────────────────────────────────────────
def login(driver):
    driver.get("https://registration.telangana.gov.in/citizen_auth.htm#!")
    for attempt in range(5):
        print(f"\n🔁 [{threading.current_thread().name}] Login Attempt {attempt + 1}")
        WebDriverWait(driver, 20).until(EC.element_to_be_clickable((By.ID, "user_type"))).click()
        WebDriverWait(driver, 10).until(
            EC.element_to_be_clickable((By.XPATH, "/html/body/div[1]/div[1]/div/form/div[1]/select/option[3]"))
        ).click()

        WebDriverWait(driver, 10).until(EC.presence_of_element_located((By.ID, "username"))).clear()
        driver.find_element(By.ID, "username").send_keys("")
        driver.find_element(By.ID, "password").clear()
        driver.find_element(By.ID, "password").send_keys("")

        captcha_text = download_and_solve_captcha(driver)
        print(f"🧠 [{threading.current_thread().name}] CAPTCHA: {captcha_text}")
        driver.find_element(By.ID, "captcha").clear()
        driver.find_element(By.ID, "captcha").send_keys(captcha_text)
        driver.find_element(By.XPATH, '//*[@id="myForm"]/div[5]/button').click()
        time.sleep(3)

        try:
            error_message = driver.find_element(By.XPATH, '//*[@id="myForm"]/h4').text.strip()
            if "Invalid Captcha" in error_message:
                print(f"❌ [{threading.current_thread().name}] CAPTCHA incorrect, retrying...")
                continue
        except:
            print(f"✅ [{threading.current_thread().name}] Login successful.")
            return True

    print(f"⚠️ [{threading.current_thread().name}] All login attempts failed.")
    return False

# ─── Open doc tab ─────────────────────────────────────────────────────────────
def open_registered_doc_tab(driver):
    WebDriverWait(driver, 20).until(
        EC.presence_of_element_located((By.XPATH, "/html/body/div[1]/div/div/form/div[6]/a"))
    ).click()
    driver.switch_to.window(driver.window_handles[-1])
    time.sleep(1.5)

# ─── Scrape one document ──────────────────────────────────────────────────────
def scrape_document(driver, doc_no, year):
    try:
        wait = WebDriverWait(driver, 10)
        wait.until(EC.element_to_be_clickable((By.ID, "deedType"))).click()
        wait.until(EC.element_to_be_clickable((By.XPATH, "/html/body/div[1]/div/select/option[1]"))).click()
        wait.until(EC.element_to_be_clickable((By.ID, "districtCode"))).click()
        district_option = wait.until(EC.element_to_be_clickable((By.XPATH, '//*[@id="districtCode"]/option[28]')))
        district_option.click()
        district_name = district_option.text.strip()

        wait.until(EC.element_to_be_clickable((By.ID, "sroCode"))).click()
        sro_option = wait.until(EC.element_to_be_clickable((By.XPATH, '//*[@id="sroCode"]/option[5]')))
        sro_option.click()
        sro_name = sro_option.text.strip()

        driver.find_element(By.ID, "doctno").clear()
        driver.find_element(By.ID, "doctno").send_keys(doc_no)
        driver.find_element(By.ID, "regyear").clear()
        driver.find_element(By.ID, "regyear").send_keys(year)

        driver.find_element(By.XPATH, '//*[@id="bean"]/div[5]/button').click()
        time.sleep(4)

        try:
            error_element = driver.find_element(By.XPATH, '//*[@id="bean"]/div[6]/div/span')
            if "No records found" in error_element.text:
                print(f"⚠️ Doc {doc_no}: No data found.")
                return None, []
        except:
            pass

        table = wait.until(EC.presence_of_element_located((By.XPATH, '//*[@id="tabAbstract"]/div[3]/table')))
        rows = table.find_elements(By.TAG_NAME, "tr")
        headers = [th.text.replace("\n", " ").strip() for th in rows[0].find_elements(By.TAG_NAME, "th")]

        records = []
        for r in rows[1:]:
            try:
                cells = r.find_elements(By.TAG_NAME, "td")
                row_data = [c.text.replace("\n", " ").strip() for c in cells]
                row_data += [doc_no, district_name, sro_name]
                if len(row_data) != len(headers) + 3:
                    raise ValueError(f"Row length mismatch: got {len(row_data)}, expected {len(headers) + 3}")
                records.append(row_data)
            except Exception as row_err:
                print(f"⚠️ Skipping malformed row in Doc {doc_no}: {row_err}")
                continue

        return headers + ["Document No", "District", "Sub-Registrar Office"], records

    except Exception as err:
        print(f"❌ Doc {doc_no}: Unexpected error - {err}")
        return None, []

# ─── Save to Excel (thread-safe) ──────────────────────────────────────────────
def save_to_excel(df):
    with excel_lock:
        if os.path.exists(excel_file):
            with pd.ExcelWriter(excel_file, mode='a', engine='openpyxl', if_sheet_exists='overlay') as writer:
                book = load_workbook(excel_file)
                sheet = book.active
                start_row = sheet.max_row
                df.to_excel(writer, index=False, header=False, startrow=start_row)
        else:
            df.to_excel(excel_file, index=False)

# ─── Worker ───────────────────────────────────────────────────────────────────
def worker(doc_queue, final_headers_holder, headers_lock, scrape_mode):
    driver = setup_driver()
    if not login(driver):
        print(f"⚠️ [{threading.current_thread().name}] Could not log in. Worker exiting.")
        driver.quit()
        return

    last_refresh = time.time()

    while True:
        try:
            doc_no = doc_queue.get_nowait()
        except Exception:
            break  # Queue empty, worker done

        # ── Session refresh every 10 min ──────────────────────────────────────
        if time.time() - last_refresh > 600:
            print(f"🔄 [{threading.current_thread().name}] Re-logging in due to timeout...")
            try:
                driver.quit()
            except:
                pass
            driver = setup_driver()
            if not login(driver):
                doc_queue.put(doc_no)
                break
            last_refresh = time.time()

        print(f"\n🔍 [{threading.current_thread().name}] Scraping Doc {doc_no}")
        try:
            open_registered_doc_tab(driver)
            hdrs, rows = scrape_document(driver, doc_no, 2025)
            driver.close()
            driver.switch_to.window(driver.window_handles[0])
        except (TimeoutException, NoSuchWindowException) as e:
            print(f"⚠️ [{threading.current_thread().name}] Window error on Doc {doc_no}: {e}")
            try:
                driver.quit()
            except:
                pass
            driver = setup_driver()
            login(driver)
            last_refresh = time.time()
            doc_queue.put(doc_no)   # Re-queue failed doc
            doc_queue.task_done()
            continue

        if rows:
            with headers_lock:
                if final_headers_holder[0] is None:
                    final_headers_holder[0] = hdrs

            df = pd.DataFrame(rows, columns=final_headers_holder[0])
            save_to_excel(df)

            # Only update progress file in range mode
            if scrape_mode == 'range':
                save_progress(doc_no)

            print(f"✅ [{threading.current_thread().name}] Doc {doc_no} saved.")
        else:
            print(f"⚠️ [{threading.current_thread().name}] No data for Doc {doc_no}.")

        doc_queue.task_done()

    try:
        driver.quit()
    except:
        pass
    print(f"🏁 [{threading.current_thread().name}] Worker done.")

# ─── Main ─────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    mode, data = read_doc_list()
    doc_queue = Queue()

    if mode == 'list':
        # ── Specific doc numbers mode ─────────────────────────────────────────
        print(f"📋 List mode: Scraping {len(data)} specific doc(s): {data}")
        for doc_no in data:
            doc_queue.put(doc_no)

    else:
        # ── Sequential range mode ─────────────────────────────────────────────
        start_doc = data
        print(f"▶️  Range mode: Starting from Doc No {start_doc}")
        for doc_no in range(start_doc, 100000):
            doc_queue.put(doc_no)

    final_headers_holder = [None]
    headers_lock = threading.Lock()

    threads = []
    for i in range(NUM_THREADS):
        t = threading.Thread(
            target=worker,
            args=(doc_queue, final_headers_holder, headers_lock, mode),
            name=f"Worker-{i+1}"
        )
        t.start()
        threads.append(t)
        time.sleep(2)   # Stagger starts to avoid simultaneous logins

    for t in threads:
        t.join()

    print("\n📁 All documents scraped successfully.")